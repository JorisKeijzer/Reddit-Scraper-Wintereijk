"""Offline tests: no network, no API key. A fake source stands in for Reddit."""

from __future__ import annotations

import openpyxl

from vocscraper import collect, db, export, label
from vocscraper.config import load_keywords, load_subreddits
from vocscraper.schema import ThreadLabel
from vocscraper.sources.common import normalize_post, walk_comment_tree

REDDIT_POST = {"kind": "t3", "data": {
    "id": "abc123", "subreddit": "PlantarFasciitis", "title": "Insoles made foot pain worse",
    "selftext": "I work 12 hour shifts on concrete.", "author": "someone", "created_utc": 1700000000.0,
    "score": 42, "num_comments": 2, "permalink": "/r/PlantarFasciitis/comments/abc123/insoles/",
}}
REDDIT_COMMENTS = [
    {"kind": "t1", "data": {"id": "c1", "parent_id": "t3_abc123", "body": "Superfeet were too hard for me.",
                            "author": "a", "created_utc": 1700000100, "score": 10, "depth": 0,
                            "replies": {"kind": "Listing", "data": {"children": [
                                {"kind": "t1", "data": {"id": "c2", "parent_id": "t1_c1", "body": "Same, try a heel cup.",
                                                        "author": "someone", "created_utc": 1700000200, "score": 3, "depth": 1,
                                                        "replies": ""}},
                                {"kind": "more", "data": {"children": ["x"]}},
                            ]}}}},
    {"kind": "t1", "data": {"id": "c3", "body": "[removed]", "author": "b", "created_utc": 1, "score": 1, "replies": ""}},
]


class FakeSource:
    name = "fake"
    supports_sitewide = True

    def search_posts(self, subreddit, query, after_utc, max_posts):
        yield normalize_post(REDDIT_POST, self.name)

    def fetch_comments(self, post_id):
        return walk_comment_tree(REDDIT_COMMENTS)


def fake_label() -> dict:
    return ThreadLabel.model_validate({
        "relevant": True, "language": "en", "thread_type": "asking_help",
        "problem_summary_nl": "Zolen maken de pijn erger bij lange diensten.", "regions": ["heel", "arch"],
        "conditions": ["plantar fasciitis"], "diagnosed_by_professional": None, "occupation": "nurse",
        "situations": ["healthcare_shift"], "duration": None, "previous_fixes": ["Superfeet"],
        "complaints_about_products": ["too_hard"], "desired_outcomes_nl": ["dienst volhouden"],
        "solutions": [
            {"solution": "Superfeet", "category": "otc_insole", "brand_or_product": "Superfeet",
             "verdict": "made_worse", "first_hand": True, "reason": "te hard", "source_id": "c1"},
            {"solution": "heel cup", "category": "heel_cup_gel_pad", "brand_or_product": None,
             "verdict": "unknown", "first_hand": False, "reason": "advies", "source_id": "c2"},
        ],
        "brands_mentioned": ["Superfeet"],
        "copper": {"mentioned": False, "sentiment": "not_mentioned", "summary_nl": ""},
        "price_mentions": [],
        "quotes": [
            {"text": "Superfeet were too hard for me.", "source_id": "c1", "kind": "fails", "nl_meaning": "Te hard."},
            {"text": "this quote is invented", "source_id": "c1", "kind": "pain", "nl_meaning": "verzonnen"},
        ],
        "emotional_intensity": 3, "medical_red_flag": False, "claims_risk": "medical",
        "marketing_insight_nl": "Hypothese: zachtere hiel.",
    }).model_dump()


def test_walk_comment_tree_flattens_and_skips_removed():
    comments = walk_comment_tree(REDDIT_COMMENTS)
    assert [c["id"] for c in comments] == ["c1", "c2"]
    assert comments[1]["depth"] == 1


def test_normalize_post_builds_full_permalink():
    post = normalize_post(REDDIT_POST, "x")
    assert post["permalink"] == "https://www.reddit.com/r/PlantarFasciitis/comments/abc123/insoles/"
    assert post["created_utc"] == 1700000000


def test_config_files_load():
    cfg = load_subreddits()
    jobs = collect.build_jobs(cfg, ["foot_core", "dutch"], [])
    assert any(j.query is None for j in jobs)  # foot_core: hele subreddit
    assert any(j.lang == "nl" and j.query == "hielspoor" for j in jobs)
    langs = {k.language for k in load_keywords()}
    assert {"en", "nl", "de"} <= langs


def test_prepare_and_import_session_labels(tmp_path, monkeypatch):
    import json

    db_path = str(tmp_path / "t.db")
    monkeypatch.setattr(collect, "load_subreddits", lambda: {
        "anchors": {"en": ["insoles"]}, "groups": {"g": {"mode": "all", "lang": "en", "subs": ["PlantarFasciitis"]}},
    })
    collect.run(db_path, FakeSource(), since=None)
    lab_dir = tmp_path / "labeling"
    [batch] = label.prepare(db_path, str(lab_dir), None)
    assert "=== THREAD abc123 ===" in batch.read_text()
    assert label.prepare(db_path, str(lab_dir), None) == []  # al klaargezet

    bad = {"post_id": "abc123", "label": {"relevant": True}}
    batch.with_suffix(".labels.json").write_text(json.dumps([bad]))
    saved, errors = label.import_labels(db_path, str(lab_dir))
    assert saved == 0 and errors and batch.exists()

    batch.with_suffix(".labels.json").write_text(json.dumps([{"post_id": "abc123", "label": fake_label()}]))
    saved, errors = label.import_labels(db_path, str(lab_dir))
    assert (saved, errors) == (1, [])
    assert (lab_dir / "imported" / batch.name).exists()


def test_collect_label_export_roundtrip(tmp_path, monkeypatch):
    db_path = str(tmp_path / "t.db")
    monkeypatch.setattr(collect, "load_subreddits", lambda: {
        "anchors": {"en": ["insoles"]}, "groups": {"g": {"mode": "anchors", "lang": "en", "subs": ["PlantarFasciitis"]}},
    })
    collect.run(db_path, FakeSource(), since=None)
    collect.run(db_path, FakeSource(), since=None)  # tweede run: geen dubbele posts

    conn = db.connect(db_path)
    assert conn.execute("SELECT COUNT(*) FROM posts").fetchone()[0] == 1
    assert conn.execute("SELECT COUNT(*) FROM comments").fetchone()[0] == 2
    assert conn.execute("SELECT author_hash FROM posts").fetchone()[0] != "someone"

    post = label.pending_posts(conn, None)[0]
    thread = label.render_thread(conn, post)
    assert "[c1]" in thread and "OP" in thread  # c2 is van de vraagsteller

    db.save_label(conn, post["id"], "test", fake_label())
    conn.commit()
    out = export.build(db_path, str(tmp_path / "out.xlsx"))
    wb = openpyxl.load_workbook(out)
    assert {"Threads", "VOC Quotes", "Oplossingen", "Wat werkt (ranking)"} <= set(wb.sheetnames)
    statuses = [r[9] for r in wb["VOC Quotes"].iter_rows(min_row=2, values_only=True)]
    assert statuses[0].startswith("Exact") and statuses[1].startswith("NIET")
    urls = [r[9] for r in wb["Oplossingen"].iter_rows(min_row=2, values_only=True)]
    assert urls[0].endswith("/abc123/insoles/c1/")
