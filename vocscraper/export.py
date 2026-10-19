"""Export labelled threads to an Excel workbook (threads, VOC quotes, solutions, rankings)."""

from __future__ import annotations

import json
import re
import sqlite3
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import db

HEADER_FILL = PatternFill("solid", fgColor="B87333")  # koperkleur
VERDICT_WEIGHT = {"works": 1.0, "partly": 0.5, "mixed": 0.0, "unknown": 0.0, "no_effect": -1.0, "made_worse": -1.5}


def _norm(text: str | None) -> str:
    return re.sub(r"\s+", " ", (text or "").replace("’", "'").replace("“", '"').replace("”", '"')).strip().lower()


def _date(ts: int | None) -> str:
    return datetime.fromtimestamp(ts or 0, tz=timezone.utc).strftime("%Y-%m-%d")


def _join(items) -> str:
    return "; ".join(str(i) for i in items or [])


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list], widths: dict[int, int] | None = None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, _ in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 18)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    return ws


def load(conn: sqlite3.Connection, include_irrelevant: bool = False):
    sql = "SELECT p.*, l.data, l.model, l.labeled_at FROM posts p JOIN labels l ON l.post_id = p.id"
    if not include_irrelevant:
        sql += " WHERE l.relevant = 1"
    for row in conn.execute(sql + " ORDER BY p.created_utc DESC"):
        comments = {c["id"]: c for c in conn.execute("SELECT * FROM comments WHERE post_id = ?", (row["id"],))}
        yield row, json.loads(row["data"]), comments


def build(db_path: str, out: str, include_irrelevant: bool = False) -> Path:
    conn = db.connect(db_path)
    wb = Workbook()
    wb.remove(wb.active)
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")

    threads, quotes, solutions, copper = [], [], [], []
    regions, conditions, situations, complaints, outcomes = Counter(), Counter(), Counter(), Counter(), Counter()
    ranking: dict[str, dict] = defaultdict(lambda: {"verdicts": Counter(), "upvotes": 0, "brands": Counter(), "first_hand": 0})
    qn = 0

    for post, lab, comments in load(conn, include_irrelevant):
        url = post["permalink"]
        source_text = {"post": _norm(f"{post['title']} {post['selftext']}")}
        source_text.update({cid: _norm(c["body"]) for cid, c in comments.items()})

        def link(source_id: str) -> str:
            return url if source_id == "post" or source_id not in comments else f"{url.rstrip('/')}/{source_id}/"

        def score(source_id: str):
            return post["score"] if source_id == "post" else (comments.get(source_id) or {"score": None})["score"]

        threads.append([
            post["id"], url, post["subreddit"], lab["language"], _date(post["created_utc"]), post["score"],
            post["num_comments"], post["title"], lab["thread_type"], lab["problem_summary_nl"],
            _join(lab["regions"]), _join(lab["conditions"]), lab["occupation"], _join(lab["situations"]),
            lab["duration"], _join(lab["previous_fixes"]), _join(lab["complaints_about_products"]),
            _join(lab["desired_outcomes_nl"]), _join(lab["brands_mentioned"]), lab["copper"]["sentiment"],
            lab["emotional_intensity"], "ja" if lab["medical_red_flag"] else "", lab["claims_risk"],
            lab["marketing_insight_nl"], "ja" if lab["relevant"] else "nee",
        ])
        for q in lab["quotes"]:
            qn += 1
            verified = _norm(q["text"]) in source_text.get(q["source_id"], "") or any(
                _norm(q["text"]) in t for t in source_text.values()
            )
            quotes.append([
                f"RS{qn:05d}", "Reddit", f"r/{post['subreddit']}", lab["problem_summary_nl"], q["text"],
                q["nl_meaning"], q["kind"], link(q["source_id"]), f"Reddit scraper {today}",
                "Exact match in bron" if verified else "NIET gevonden in bron - handmatig checken",
                _join(lab["regions"]), lab["occupation"] or "", score(q["source_id"]), lab["language"],
            ])
        for s in lab["solutions"]:
            up = score(s["source_id"]) or 0
            solutions.append([
                post["id"], f"r/{post['subreddit']}", s["category"], s["solution"], s["brand_or_product"],
                s["verdict"], "ja" if s["first_hand"] else "nee", s["reason"], up, link(s["source_id"]),
                _join(lab["regions"]), _join(lab["conditions"]),
            ])
            r = ranking[s["category"]]
            r["verdicts"][s["verdict"]] += 1
            r["upvotes"] += max(up, 0)
            r["first_hand"] += int(s["first_hand"])
            if s["brand_or_product"]:
                r["brands"][s["brand_or_product"]] += 1
        if lab["copper"]["mentioned"]:
            copper.append([post["id"], url, f"r/{post['subreddit']}", lab["copper"]["sentiment"], lab["copper"]["summary_nl"], post["title"]])
        regions.update(lab["regions"])
        conditions.update(c.lower() for c in lab["conditions"])
        situations.update(lab["situations"])
        complaints.update(lab["complaints_about_products"])
        outcomes.update(lab["desired_outcomes_nl"])

    _sheet(wb, "Threads", [
        "Post ID", "URL", "Subreddit", "Taal", "Datum", "Score", "Reacties", "Titel", "Type", "Probleem (NL)",
        "Regio", "Aandoeningen", "Beroep", "Situaties", "Duur", "Eerder geprobeerd", "Klachten over producten",
        "Gewenste uitkomst", "Merken", "Koper-sentiment", "Emotie 1-5", "Medische red flag", "Claimrisico",
        "Marketing-inzicht (hypothese)", "Relevant",
    ], threads, {2: 30, 8: 40, 10: 50, 24: 50})
    _sheet(wb, "VOC Quotes", [
        "Source ID", "Platform", "Community", "Problem context (NL)", "Exact excerpt (Q)", "NL meaning (T)",
        "Soort", "Original URL", "Provenance", "Verification status", "Anatomy tag", "Occupation tag",
        "Upvotes", "Taal",
    ], quotes, {4: 40, 5: 60, 6: 50, 8: 35, 10: 25})
    _sheet(wb, "Oplossingen", [
        "Post ID", "Community", "Categorie", "Oplossing", "Merk/product", "Oordeel", "Eigen ervaring",
        "Waarom", "Upvotes", "Bron", "Regio", "Aandoeningen",
    ], sorted(solutions, key=lambda r: -(r[8] or 0)), {4: 40, 8: 50, 10: 35})

    rank_rows = []
    for cat, r in ranking.items():
        v, total = r["verdicts"], sum(r["verdicts"].values())
        net = sum(VERDICT_WEIGHT[k] * n for k, n in v.items()) / total
        rank_rows.append([
            cat, total, r["first_hand"], v["works"], v["partly"], v["no_effect"], v["made_worse"], v["mixed"] + v["unknown"],
            round(net, 2), r["upvotes"], _join(f"{b} ({n})" for b, n in r["brands"].most_common(8)),
        ])
    _sheet(wb, "Wat werkt (ranking)", [
        "Categorie", "Vermeldingen", "Eigen ervaring", "Werkt", "Deels", "Geen effect", "Erger", "Gemengd/onbekend",
        "Netto-score (-1.5..1)", "Upvotes totaal", "Top merken",
    ], sorted(rank_rows, key=lambda r: -r[1]), {1: 26, 11: 60})

    problem_rows = []
    for label, counter in [("Regio", regions), ("Aandoening", conditions), ("Situatie", situations),
                           ("Klacht over product", complaints), ("Gewenste uitkomst", outcomes)]:
        problem_rows += [[label, k, n] for k, n in counter.most_common(60)]
    _sheet(wb, "Problemen (telling)", ["Dimensie", "Waarde", "Aantal threads"], problem_rows, {2: 60})
    _sheet(wb, "Koper & wellness", ["Post ID", "URL", "Community", "Sentiment", "Wat wordt er gezegd", "Titel"], copper, {2: 30, 5: 70, 6: 40})

    coverage = conn.execute(
        """SELECT p.subreddit, p.lang_hint, COUNT(*) AS posts, COUNT(l.post_id) AS gelabeld,
                  SUM(COALESCE(l.relevant, 0)) AS relevant
           FROM posts p LEFT JOIN labels l ON l.post_id = p.id GROUP BY p.subreddit ORDER BY posts DESC"""
    ).fetchall()
    _sheet(wb, "Dekking", ["Subreddit", "Taal", "Posts", "Gelabeld", "Relevant"], [list(r) for r in coverage])

    out_path = Path(out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
